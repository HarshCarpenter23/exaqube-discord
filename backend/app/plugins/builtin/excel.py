"""excel plugin — build a formatted .xlsx from one or more result tables.

One sheet per input table, with a bold frozen header row, an autofilter, and
sensible column widths — a real workbook, not a CSV with an .xlsx extension.
Bounded by artifact_max_rows so "export every message to Excel" is rejected
rather than exhausting memory.

Consumes one or more `table` handles via input_refs.
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.plugins.base import Plugin, PluginContext, PluginError, PluginResult
from app.plugins.registry import register
from app.serialization import jsonable

_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_MAX_WIDTH = 60


@register
class ExcelPlugin(Plugin):
    name = "excel"
    description = (
        "Build a formatted Excel workbook (.xlsx) from one or more result tables. "
        "Pass input_refs with the table handles; give sheet_names to label them. "
        "Returns a downloadable artifact."
    )
    input_schema = {
        "type": "object",
        "properties": {
            "input_refs": {
                "type": "array",
                "items": {"type": "string"},
                "minItems": 1,
                "description": "Handles of the tables to include, one per sheet.",
            },
            "sheet_names": {"type": "array", "items": {"type": "string"}},
            "filename": {"type": "string"},
        },
        "required": ["input_refs"],
        "additionalProperties": False,
    }
    consumes = {"table"}

    async def execute(self, args: dict, ctx: PluginContext) -> PluginResult:
        tables = ctx.inputs["input_refs"]                 # list[PluginResult], kind-checked by loop
        sheet_names = args.get("sheet_names") or []

        total_rows = sum(t.data["row_count"] for t in tables)
        if total_rows > ctx.artifact_max_rows:
            raise PluginError(
                "too_large",
                f"Export is {total_rows} rows, over the {ctx.artifact_max_rows} limit. "
                "Narrow the query first.",
                retryable=True,
            )

        await ctx.emit("tool_progress", {"message": f"building workbook ({len(tables)} sheet(s))"})
        workbook = self._build_workbook(tables, sheet_names)

        buffer = BytesIO()
        workbook.save(buffer)
        data = buffer.getvalue()
        if len(data) > ctx.artifact_max_bytes:
            raise PluginError("too_large", "Workbook exceeds the size limit.", retryable=True)

        filename = args.get("filename") or "export.xlsx"
        if not filename.endswith(".xlsx"):
            filename += ".xlsx"
        info = ctx.artifacts.write_bytes(data, filename, _XLSX_MIME)
        return PluginResult(kind="artifact", data=info.as_dict(), meta={})

    def _build_workbook(self, tables, sheet_names) -> Workbook:
        workbook = Workbook()
        workbook.remove(workbook.active)  # drop the default empty sheet
        for i, table in enumerate(tables):
            title = sheet_names[i] if i < len(sheet_names) else f"Sheet{i + 1}"
            self._write_sheet(workbook, title[:31], table.data["columns"], table.data["rows"])
        return workbook

    def _write_sheet(self, workbook: Workbook, title: str, columns: list[str], rows: list[list]):
        sheet = workbook.create_sheet(title=title)
        sheet.append(columns)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in rows:
            sheet.append([jsonable(v) for v in row])

        sheet.freeze_panes = "A2"
        if columns:
            sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
        self._fit_columns(sheet, columns, rows)

    def _fit_columns(self, sheet, columns: list[str], rows: list[list]) -> None:
        for idx, name in enumerate(columns):
            widest = len(str(name))
            for row in rows[:200]:  # sample for width, don't scan millions
                if idx < len(row) and row[idx] is not None:
                    widest = max(widest, len(str(jsonable(row[idx]))))
            sheet.column_dimensions[get_column_letter(idx + 1)].width = min(widest + 2, _MAX_WIDTH)
