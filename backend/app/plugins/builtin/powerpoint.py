"""powerpoint plugin — build a .pptx deck from tables, charts, and prose.

A title slide plus one content slide per referenced result: a chart_spec becomes
a native PowerPoint chart, a table becomes a table shape (bounded), and the
agent-written prose is placed alongside. Bounded by pptx_max_slides.

Consumes `table` and `chart_spec` handles via input_refs. Uses python-pptx's
native charts, so no image-rendering dependency is needed.
"""

from __future__ import annotations

from io import BytesIO

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

from app.plugins.base import Plugin, PluginContext, PluginError, PluginResult
from app.plugins.registry import register
from app.serialization import jsonable

_PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
_TABLE_ROW_LIMIT = 12       # rows shown per table slide
_CHART_TYPE = {"line": XL_CHART_TYPE.LINE, "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
               "distribution": XL_CHART_TYPE.COLUMN_CLUSTERED}


@register
class PowerPointPlugin(Plugin):
    name = "powerpoint"
    description = (
        "Build a PowerPoint deck (.pptx) from result tables and charts plus prose you write. "
        "Pass input_refs (table or chart handles) — each becomes a slide — with a title and "
        "optional per-slide headings and notes. Returns a downloadable artifact."
    )
    input_schema = {
        "type": "object",
        "properties": {
            "title": {"type": "string"},
            "summary": {"type": "string", "description": "Intro prose for the title slide."},
            "input_refs": {"type": "array", "items": {"type": "string"}},
            "slide_headings": {"type": "array", "items": {"type": "string"}},
            "slide_notes": {"type": "array", "items": {"type": "string"}},
            "filename": {"type": "string"},
        },
        "required": ["title", "input_refs"],
        "additionalProperties": False,
    }
    consumes = {"table", "chart_spec"}

    async def execute(self, args: dict, ctx: PluginContext) -> PluginResult:
        results = ctx.inputs["input_refs"]              # list[PluginResult], kind-checked by loop
        if len(results) + 1 > ctx.pptx_max_slides:
            raise PluginError(
                "too_large",
                f"Deck would have {len(results) + 1} slides, over the {ctx.pptx_max_slides} limit.",
                retryable=True,
            )

        headings = args.get("slide_headings") or []
        notes = args.get("slide_notes") or []

        await ctx.emit("tool_progress", {"message": f"building deck ({len(results) + 1} slides)"})
        prs = Presentation()
        self._title_slide(prs, args["title"], args.get("summary", ""))
        for i, result in enumerate(results):
            heading = headings[i] if i < len(headings) else f"Slide {i + 1}"
            note = notes[i] if i < len(notes) else ""
            self._content_slide(prs, heading, note, result)

        buffer = BytesIO()
        prs.save(buffer)
        data = buffer.getvalue()
        if len(data) > ctx.artifact_max_bytes:
            raise PluginError("too_large", "Deck exceeds the size limit.", retryable=True)

        filename = args.get("filename") or "deck.pptx"
        if not filename.endswith(".pptx"):
            filename += ".pptx"
        info = ctx.artifacts.write_bytes(data, filename, _PPTX_MIME)
        return PluginResult(kind="artifact", data=info.as_dict(), meta={})

    def _title_slide(self, prs: Presentation, title: str, summary: str) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[0])   # Title Slide layout
        slide.shapes.title.text = title
        if len(slide.placeholders) > 1:
            slide.placeholders[1].text = summary

    def _content_slide(self, prs: Presentation, heading: str, note: str, result: PluginResult) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[5])   # Title Only layout
        slide.shapes.title.text = heading
        if result.kind == "chart_spec":
            self._add_chart(slide, result.data)
        else:
            self._add_table(slide, result.data)
        if note:
            self._add_note(slide, note)

    def _add_chart(self, slide, spec: dict) -> None:
        columns, rows = spec["columns"], spec["rows"]
        x, y = spec["x"], spec.get("y")
        if not y:
            return
        xi, yi = columns.index(x), columns.index(y)
        chart_data = CategoryChartData()
        chart_data.categories = [str(jsonable(r[xi])) for r in rows]
        chart_data.add_series(y, [r[yi] for r in rows])
        chart_type = _CHART_TYPE.get(spec.get("type"), XL_CHART_TYPE.LINE)
        slide.shapes.add_chart(chart_type, Inches(0.7), Inches(1.5), Inches(8.6), Inches(4.8), chart_data)

    def _add_table(self, slide, data: dict) -> None:
        columns = data["columns"]
        rows = data["rows"][:_TABLE_ROW_LIMIT]
        table_shape = slide.shapes.add_table(
            len(rows) + 1, len(columns), Inches(0.7), Inches(1.5), Inches(8.6), Inches(0.4)
        )
        table = table_shape.table
        for c, name in enumerate(columns):
            table.cell(0, c).text = str(name)
        for r, row in enumerate(rows, start=1):
            for c, value in enumerate(row):
                table.cell(r, c).text = "" if value is None else str(jsonable(value))

    def _add_note(self, slide, note: str) -> None:
        box = slide.shapes.add_textbox(Inches(0.7), Inches(6.4), Inches(8.6), Inches(0.8))
        frame = box.text_frame
        frame.word_wrap = True
        frame.text = note
        frame.paragraphs[0].font.size = Pt(12)
